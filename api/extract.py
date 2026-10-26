#!/usr/bin/env python3
"""Extract plain text from PDF / DOCX files for the Local AI knowledge base.

Reads a file path and writes UTF-8 text to stdout. Errors go to stderr with a
non-zero exit code. No network access, no data leaves the machine.
"""
import os
import sys


def extract_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts)


def extract_docx(path):
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def main():
    if len(sys.argv) < 2:
        sys.stderr.write("usage: extract.py <file>\n")
        return 2
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        text = extract_pdf(path)
    elif ext == ".docx":
        text = extract_docx(path)
    else:
        with open(path, "r", encoding="utf-8", errors="replace") as fh:
            text = fh.read()
    sys.stdout.buffer.write(text.encode("utf-8", errors="replace"))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001
        sys.stderr.write("ERROR: %s: %s\n" % (type(exc).__name__, exc))
        sys.exit(1)
